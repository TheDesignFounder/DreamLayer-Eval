"""Read-only historical recovery. Filename matches are candidates, never identities."""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from .contracts import file_hash, redact
from .media import probe
from .types import JsonObject, PathLike


def png_record(path: PathLike, root: PathLike) -> JsonObject:
    "Recover image hashes and embedded metadata without changing the source file."
    data = {
        "source": str(path.relative_to(root)),
        "original_filename": path.name,
        "sha256": file_hash(path),
        "bytes": path.stat().st_size,
    }
    with Image.open(path) as image:
        data["dimensions"] = list(image.size)
        data["png_metadata"] = redact(dict(image.info))
    graph = data["png_metadata"].get("prompt")
    try:
        graph = json.loads(graph) if isinstance(graph, str) else graph
        nodes = list(graph.values()) if isinstance(graph, dict) else []
        texts = [
            n["inputs"]["prompt"]
            for n in nodes
            if isinstance(n, dict) and isinstance(n.get("inputs", {}).get("prompt"), str)
        ]
        if not texts:
            texts = [
                n["inputs"]["text"]
                for n in nodes
                if isinstance(n, dict)
                and n.get("class_type") == "CLIPTextEncode"
                and isinstance(n.get("inputs", {}).get("text"), str)
            ]
        data["embedded_prompts"] = texts
        data["node_types"] = sorted(
            {n.get("class_type", "unknown") for n in nodes if isinstance(n, dict)}
        )
    except (ValueError, TypeError):
        data.update(embedded_prompts=[], node_types=[], metadata_error="unparseable graph")
    data["actual_endpoint_evidence"] = (
        "StabilityStableImageUltraNode"
        if "StabilityStableImageUltraNode" in data.get("node_types", [])
        else None
    )
    data["historical_model_identity"] = "unknown beyond embedded workflow evidence"
    return data


def workbook_records(path: PathLike, root: PathLike) -> tuple[list[JsonObject], list[JsonObject]]:
    "Read workbook observations with sheet and source-file provenance."
    workbook = load_workbook(path, read_only=True, data_only=False)
    rows, sheets = [], []
    workbook_hash = file_hash(path)
    for sheet in workbook:
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, ())
        count = 0
        for index, values in enumerate(iterator, 2):
            row = {
                str(k): v.isoformat() if hasattr(v, "isoformat") else v
                for k, v in zip(header, values)
                if k is not None and v is not None
            }
            if not row.get("run_id"):
                continue
            filename = Path(str(row.get("asset_value", row.get("full_path", "")))).name
            if not filename.endswith(".png"):
                try:
                    filename = json.loads(row.get("generated_images", "[]"))[0]
                except (ValueError, IndexError, TypeError):
                    filename = ""
            rows.append(
                {
                    "workbook": str(path.relative_to(root)),
                    "workbook_sha256": workbook_hash,
                    "sheet": sheet.title,
                    "row": index,
                    "original": redact(row),
                    "filename_candidate": filename,
                }
            )
            count += 1
        sheets.append({"workbook": path.name, "sheet": sheet.title, "run_rows": count})
    workbook.close()
    return rows, sheets


def human_rankings(path: PathLike) -> JsonObject:
    "Read the supported historical CSV while preserving original rank and tie fields."
    with open(path, newline="", encoding="utf-8-sig") as f:
        raw = list(csv.DictReader(f))
    rows = []
    for i, row in enumerate(raw):
        label = row.get("Quality Ranking", "")
        rows.append(
            {
                "source_row": i + 2,
                "original": row,
                "display_order": i + 1,
                "tie_group": "explicit-even" if "(Even)" in label else None,
                "numeric_score": None,
                "asset_id": None,
                "link_status": "unresolved_model_label_only",
            }
        )
    return {
        "source_filename": Path(path).name,
        "source_sha256": file_hash(path),
        "test": "Test 1",
        "author": "historical-author-1",
        "exploratory": True,
        "blinded": False,
        "vlm_used": False,
        "independent_rater_count": 1,
        "rankings": rows,
    }


def recover(downloads: PathLike, output: PathLike, limit: int | None = None) -> JsonObject:
    "Inventory local historical evidence and write an explicit reconciliation report."
    root, output = Path(downloads), Path(output)
    output.mkdir(parents=True, exist_ok=True)
    paths = sorted(
        p
        for name in ("DreamLayer Data", "DreamLayer Data 2")
        for p in (root / name / "Image Data Generated").rglob("*.png")
    )
    if limit:
        paths = paths[:limit]
    images = [png_record(p, root) for p in paths]
    byname, byhash = defaultdict(list), defaultdict(list)
    for im in images:
        byname[im["original_filename"]].append(im)
        byhash[im["sha256"]].append(im["source"])
    records, sheets = [], []
    for p in sorted((root / "DreamLayer Data 2").glob("*.xlsx")):
        rr, ss = workbook_records(p, root)
        records.extend(rr)
        sheets.extend(ss)
    conflicts, unresolved = [], []
    for row in records:
        prompt = str(row["original"].get("prompt", row["original"].get("prompt_text", "")))
        candidates = byname[row["filename_candidate"]]
        matches = [
            im
            for im in candidates
            if any(
                " ".join(prompt.split()) == " ".join(p.split())
                for p in im.get("embedded_prompts", [])
            )
        ]
        row["candidate_hashes"] = sorted({im["sha256"] for im in candidates})
        row["verified_hashes"] = sorted({im["sha256"] for im in matches})
        row["link_status"] = (
            "embedded_prompt_and_filename_agree"
            if len(row["verified_hashes"]) == 1
            else "unresolved"
        )
        row["fid_validity"] = (
            "invalid_legacy_result" if row["original"].get("fid_score") is not None else "missing"
        )
        row["parameter_enforcement"] = "unverified"
        if candidates and not matches and any(im.get("embedded_prompts") for im in candidates):
            conflicts.append(
                {
                    "workbook": row["workbook"],
                    "sheet": row["sheet"],
                    "row": row["row"],
                    "filename": row["filename_candidate"],
                    "reason": "Filename candidate has a different embedded prompt",
                }
            )
        if row["link_status"] == "unresolved":
            unresolved.append(
                {
                    "workbook": row["workbook"],
                    "sheet": row["sheet"],
                    "row": row["row"],
                    "reason": "No unique independently supported content identity",
                }
            )
    video_paths = sorted(
        p
        for p in root.glob("*.mp4")
        if p.name.startswith(("video_", "marketing_ugc", "voice_clone"))
    )
    if limit:
        video_paths = video_paths[:limit]

    def video(p):
        try:
            return {
                "source": p.name,
                "sha256": file_hash(p),
                "media": probe(p),
                "parent_generation": None,
            }
        except Exception as exc:
            return {"source": p.name, "error": str(exc)}

    with ThreadPoolExecutor(max_workers=4) as pool:
        videos = list(pool.map(video, video_paths))
    # Candidate workflows are retained as evidence, never asserted to be a clip's
    # generating workflow from a shared filename or model label alone.
    from .contracts import redact

    workflows = []
    for path in sorted(root.glob("*.json")):
        if not path.name.startswith(
            ("comfy_", "video_", "voice_clone", "11_voice_clone", "marketing_ugc")
        ):
            continue
        try:
            raw = json.loads(path.read_text())
            if not isinstance(raw, dict) or not (
                "nodes" in raw
                or any(isinstance(v, dict) and "class_type" in v for v in raw.values())
            ):
                continue
            workflows.append(
                {
                    "source": path.name,
                    "source_sha256": file_hash(path),
                    "workflow": redact(raw),
                    "link_status": "candidate_only_no_verified_generation_link",
                    "asset_id": None,
                }
            )
        except (ValueError, OSError):
            continue
    human_path = root / "Human Video Eval - Sheet1.csv"
    human = human_rankings(human_path) if human_path.exists() else None
    summary = {
        "schema_version": 1,
        "image_files": len(images),
        "unique_image_hashes": len(byhash),
        "images_with_embedded_prompts": sum(bool(im.get("embedded_prompts")) for im in images),
        "video_files": len(videos),
        "unique_video_hashes": len({v["sha256"] for v in videos if "sha256" in v}),
        "workflow_candidates": len(workflows),
        "workbook_rows": len(records),
        "workbook_sheets": sheets,
        "verified_row_links": sum(r["link_status"] != "unresolved" for r in records),
        "conflicts": conflicts,
        "unresolved": unresolved,
        "duplicate_image_content": {k: v for k, v in byhash.items() if len(v) > 1},
        "missingness": {
            m: sum(r["original"].get(m) in (None, "") for r in records)
            for m in ("clip_score_mean", "fid_score", "macro_f1")
        },
        "notes": [
            "Overlapping workbook rows are not independent trials",
            "Different hashes do not establish different source generations",
            "Historical reference statistics cache is unavailable",
            "Stability Ultra workflow evidence conflicts with SD Turbo labels; no historical relabeling performed",
        ],
        "human": human,
    }
    (output / "reconciliation.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False)
    )
    for name, data in (
        ("images", images),
        ("workbook_rows", records),
        ("videos", videos),
        ("workflows", workflows),
    ):
        with (output / (name + ".jsonl")).open("w") as f:
            for row in data:
                f.write(json.dumps(row, ensure_ascii=False, allow_nan=False) + "\n")
    return summary
